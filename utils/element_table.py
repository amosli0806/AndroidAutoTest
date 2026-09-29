# utils/element_table.py
"""应用元素库的表格（Excel）导入导出。

格式约定（**一个应用一张工作表**，表名就是应用名）：

    所属模块 | 名称       | 定位方式 | 定位值                              | 备注
    主图态    | 放大       | 资源ID   | com.baidu.naviauto:id/zoom_in_btn  |
    检索      | SUG联想结果 | XPath   | //*[@resource-id="..."]/android... |

要点：
1. 工作表 = 所属应用，所以表里**不再重复放一列应用名**；要换应用就把整行挪到另一张表。
2. **不带 id 列**：id 是内部主键，对用户没意义。导入时用
   「所属应用(表名) + 所属模块 + 名称」当业务主键 —— 命中现有元素就覆盖它的
   定位方式/定位值/备注并**保留原 id**，没命中才新增。保留 id 很关键：步骤是通过
   element_id 引用元素的，重建 id 会让引用失效（执行时报「元素ID 已失效」）。
3. **导入只新增 + 覆盖，不删除**：表格里少了某行不代表要删元素，误删会静默打断
   步骤引用，风险不对称。要删在界面里删（那边有引用检查）。

openpyxl 延迟导入：主程序启动路径对时间敏感，不为一个按需功能拖慢启动。
"""
import re
from typing import List, Tuple

from models.element_model import LOC_TYPES

# 表头（导入时认这一行就跳过；没有表头也能导）
HEADERS = ("所属模块", "名称", "定位方式", "定位值", "备注")

# Excel 工作表名不允许的字符
_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")

# 单元格文本长度上限（Excel 单格上限 32767，这里收紧到够用又不至于爆）
MAX_CELL_CHARS = 2000

# 写表时按行数估的行高（Excel 不会自动撑开换行单元格）
ROW_HEIGHT_PER_LINE = 15

# 列宽：定位值（资源 ID / XPath）最长，给足
COLUMN_WIDTHS = (20, 26, 12, 56, 24)


# ============================================================
# 导出方向：元素 -> 一行五列
# ============================================================
def element_to_row(element) -> List[str]:
    """一个元素 -> 一行 [所属模块, 名称, 定位方式, 定位值, 备注]。"""
    return [element.module or "", element.name or "",
            element.loc_type or "", element.loc_value or "",
            element.remark or ""]


def group_by_app(elements) -> List[Tuple[str, list]]:
    """按所属应用分组（保持首次出现的顺序），返回 [(应用名, [元素, ...]), ...]。

    空应用归到「未指定应用」—— 否则那批元素在导出时会被漏掉。
    """
    groups = {}
    order = []
    for element in elements:
        app = (element.app or "").strip() or "未指定应用"
        if app not in groups:
            groups[app] = []
            order.append(app)
        groups[app].append(element)
    return [(app, groups[app]) for app in order]


def app_to_sheet_rows(elements) -> List[list]:
    """一个应用下的全部元素 -> 工作表数据行（表头由 write_xlsx 加）。"""
    return [element_to_row(e) for e in elements]


# ============================================================
# 导入方向：一行五列 -> 元素字段
# ============================================================
def parse_sheet(rows) -> Tuple[List[Tuple[str, str, str, str, str]], List[str]]:
    """一个工作表的数据行 -> ([(模块, 名称, 定位方式, 定位值, 备注)], 警告列表)。

    逐行校验：模块/名称/定位值不能为空，定位方式必须是 LOC_TYPES 里的一种；
    不合格的行跳过并记一条带行号的警告，不整表失败。
    """
    parsed: List[Tuple[str, str, str, str, str]] = []
    warnings: List[str] = []
    header_pending = True
    for line_no, row in enumerate(rows, start=1):
        cells = [(row[i] if len(row) > i else "") or "" for i in range(5)]
        module, name, loc_type, loc_value, remark = (str(c).strip() for c in cells)
        if header_pending and name == HEADERS[1]:
            header_pending = False
            continue
        header_pending = False
        if not any((module, name, loc_type, loc_value, remark)):
            continue
        if not name:
            warnings.append(f"第 {line_no} 行没有元素名称，已跳过")
            continue
        if not module:
            warnings.append(f"第 {line_no} 行「{name}」没有所属模块，已跳过")
            continue
        if not loc_value:
            warnings.append(f"第 {line_no} 行「{name}」没有定位值，已跳过")
            continue
        if loc_type not in LOC_TYPES:
            warnings.append(
                f"第 {line_no} 行「{name}」的定位方式「{loc_type or '空'}」不是 "
                f"{'/'.join(LOC_TYPES)} 之一，已跳过")
            continue
        parsed.append((module, name, loc_type, loc_value[:MAX_CELL_CHARS],
                       remark[:MAX_CELL_CHARS]))
    return parsed, warnings


# ============================================================
# Excel 读写
# ============================================================
def _safe_sheet_name(name: str, used: set) -> str:
    """工作表名合法化：去非法字符、截到 31 字、重名补序号。"""
    clean = _INVALID_SHEET_CHARS.sub("_", (name or "").strip()) or "未指定应用"
    clean = clean[:31]
    candidate, index = clean, 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = clean[:31 - len(suffix)] + suffix
        index += 1
    used.add(candidate)
    return candidate


def write_xlsx(path: str, sheets) -> int:
    """写出 xlsx。sheets = [(应用名, 该应用元素的数据行), ...]，一个应用一张表。

    返回写出的工作表数。表名按 Excel 规则合法化并去重。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    workbook.remove(workbook.active)   # 去掉默认的空表
    used = set()
    wrap = Alignment(wrap_text=True, vertical="top")
    for sheet_name, rows in sheets:
        sheet = workbook.create_sheet(_safe_sheet_name(sheet_name, used))
        sheet.append(list(HEADERS))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            sheet.append([str(c)[:MAX_CELL_CHARS] for c in row])
            for cell in sheet[sheet.max_row]:
                cell.alignment = wrap
            # 行高按行数估（Excel 不会自动撑开换行单元格）
            line_count = max((len(str(c).splitlines()) for c in row), default=1)
            sheet.row_dimensions[sheet.max_row].height = \
                max(1, line_count) * ROW_HEIGHT_PER_LINE
        # 定位方式做成下拉：手写「资源id」这种错字会被导入校验拦下，不如根本不让写错
        validation = DataValidation(
            type="list", formula1='"' + ",".join(LOC_TYPES) + '"', allow_blank=True)
        validation.error = f"定位方式必须是 {'/'.join(LOC_TYPES)} 之一"
        validation.errorTitle = "定位方式不合法"
        sheet.add_data_validation(validation)
        validation.add(f"C2:C{max(sheet.max_row, 2)}")
        for index, width in enumerate(COLUMN_WIDTHS, start=1):
            sheet.column_dimensions[
                sheet.cell(row=1, column=index).column_letter].width = width
    workbook.save(path)
    return len(workbook.sheetnames)


def read_xlsx(path: str) -> List[Tuple[str, List[list]]]:
    """读出 xlsx：返回 [(工作表名, 数据行), ...]，每行已裁成五列。

    表头不在这里剔 —— 统一由 parse_sheet 负责（见那里的注释）。
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    result = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for raw in sheet.iter_rows(values_only=True):
                cells = list(raw[:5]) if raw else []
                while len(cells) < 5:
                    cells.append(None)
                rows.append(["" if c is None else str(c) for c in cells])
            result.append((sheet.title, rows))
    finally:
        workbook.close()
    return result
