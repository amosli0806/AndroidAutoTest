# utils/test_data.py
"""测试数据读取：CSV / XLSX → [{表头: 值}, ...]，供数据驱动执行。

约定（与界面口径一致）：
  * 第一行是表头，**表头 = 变量名**；步骤里用 `${列名}` 引用。
  * 其余每行是一组数据，每组把用例跑一遍。
  * 单元格一律按字符串处理（"01" 不会变成 1），空单元格 → 空串。
  * 编码：UTF-8（含 BOM）优先，失败回退 GBK。

openpyxl 延迟导入：只在读 xlsx 时才 import（主程序启动路径对时间敏感，
参考 utils/element_table.py 的做法），所以不读 Excel 时零开销。
"""
import csv
import io
import os
from typing import Dict, List


def _read_text(path: str) -> str:
    with open(path, 'rb') as f:
        data = f.read()
    for enc in ('utf-8-sig', 'gbk'):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别的文件编码（支持 UTF-8 / GBK）")


def _rows_to_dicts(rows: List[list]) -> List[Dict[str, str]]:
    """把 [表头行, 数据行...] 归一成 [{列名: 值}]，并做空表/空表头/重复表头校验。"""
    # 去掉全空行（Excel 表格底部常见一堆空行）
    rows = [r for r in rows if any((c or "") != "" for c in r)]
    if not rows:
        raise ValueError("数据文件是空的")

    headers = [str(h).strip() for h in rows[0]]
    if not any(headers):
        raise ValueError("第一行（表头）是空的")
    if "" in headers:
        raise ValueError("表头里有空的列名（每个表头都要填）")
    if len(set(headers)) != len(headers):
        dup = sorted({h for h in headers if headers.count(h) > 1})
        raise ValueError("表头有重复列名：" + "、".join(dup))

    result = []
    for row in rows[1:]:
        result.append({
            headers[j]: ("" if j >= len(row) else str(row[j]).strip())
            for j in range(len(headers))
        })
    return result


def read_csv(path: str) -> List[Dict[str, str]]:
    text = _read_text(path)
    rows = list(csv.reader(io.StringIO(text)))
    return _rows_to_dicts(rows)


def read_xlsx(path: str) -> List[Dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ValueError("读取 .xlsx 需要 openpyxl，当前环境未安装") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = [[("" if c is None else str(c)) for c in row]
                for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return _rows_to_dicts(rows)


def read_file(path: str) -> List[Dict[str, str]]:
    """按扩展名分发。返回 [{表头: 值}, ...]（不含表头本身）。"""
    ext = os.path.splitext(str(path))[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        return read_xlsx(path)
    if ext in ('.csv', '.txt'):
        return read_csv(path)
    raise ValueError(f"不支持的数据文件格式「{ext}」（支持 .csv / .xlsx）")
