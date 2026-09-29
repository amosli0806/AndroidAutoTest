# utils/voice_table.py
"""语音用例的表格（Excel）导入导出。

格式约定（一个分组一张工作表）：

    用例名称              | 操作步骤              | 预期结果
    打开地图 - 地图在全屏  | 1. 你好本田           | 1. 主驾，在呢，我在
                         | 2. 我要打开导航页面     | 2. 正在为您切换为全屏地图
                         | 3. 我要打开导航页面     | 3. 当前已为全屏地图
    打开地图 - 地图在半屏  | 1. 你好本田           | 1. 主驾，在呢，我在
                         | 2. 我要打开导航页面     |

要点：
1. **工作表 = 分组**：表名就是分组名，一张表装该分组下的全部用例；一个用例占一行。
2. **操作步骤 / 预期结果两列里是「换行 + 序号」的多行文本**，一个单元格装完该用例的
   全部步骤 / 全部预期。
3. **序号就是「步骤号」**：第 N 条预期前面的数字，指的是它对应的那条操作步骤。
   内部模型里检测步骤归属于紧挨在它前面的那句播报，所以「一句播报配多条预期」是成立的，
   而「有预期却没有对应步骤」不成立。用步骤号当预期序号，两边数量不等时也不会错位：
   某条步骤没有预期，它就不出现在预期列里。
4. 同一句播报配了多条预期时，预期列会出现多行**相同序号**——导入时它们各自还原成
   一条独立的检测步骤。刻意不合并成一条：合并会把「各自都要通过」悄悄放宽成
   「命中任一即通过」。

openpyxl 延迟导入：主程序启动路径对时间敏感，不为一个按需功能拖慢启动。
"""
import re
from typing import List, Tuple

from models.voice_model import (
    DEFAULT_DELAY, KIND_PHRASE, KIND_VERIFY, VoiceStep, _clean_keywords)

# 表头（导入时认这一行就跳过；没有表头也能导）
HEADERS = ("用例名称", "操作步骤", "预期结果")

# 关键词分隔：与界面上检测步骤输入框的规则完全一致（中英文逗号都认）
_KEYWORD_SPLIT = re.compile(r"[,，]")

# 行首序号：「1. xxx」「2、xxx」「3) xxx」「4：xxx」都认
_STEP_NO_RE = re.compile(r"^\s*(\d+)\s*[.、．)）:：]\s*(.*)$")

# Excel 工作表名不允许的字符
_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")

# 单元格文本长度上限（Excel 单格上限 32767，这里收紧到够用又不至于爆）
MAX_CELL_CHARS = 2000

# 写表时按行数估的行高（Excel 不会自动撑开换行单元格，给个够用的高度看着才像样）
ROW_HEIGHT_PER_LINE = 15


# ============================================================
# 导出方向：用例 -> 一行三列
# ============================================================
def case_to_row(case) -> List[str]:
    """一个用例 -> 一行 [用例名, 操作步骤文本, 预期结果文本]。"""
    step_lines: List[str] = []
    expect_lines: List[str] = []
    step_no = 0
    last_step_no = None
    for step in case.steps:
        if step.kind == KIND_PHRASE:
            step_no += 1
            last_step_no = step_no
            step_lines.append(f"{step_no}. {step.text or ''}")
            continue
        # 检测步骤：归属于紧挨在它前面的那句播报，序号就用那句播报的步骤号
        if last_step_no is None:
            continue      # 开头就是检测（异常数据）：没有归属，丢掉
        expect_lines.append(f"{last_step_no}. {'，'.join(step.keywords or [])}")
    return [case.name, "\n".join(step_lines), "\n".join(expect_lines)]


def group_to_sheet_rows(cases) -> List[list]:
    """一个分组的全部用例 -> 工作表的数据行（表头由 write_xlsx 加）。"""
    return [case_to_row(case) for case in cases]


# ============================================================
# 导入方向：一行三列 -> 用例
# ============================================================
def _parse_numbered(text) -> List[Tuple[int, str]]:
    """把「1. xxx\\n2. yyy」解析成 [(序号, 内容)]。

    没写序号的行按出现顺序自动补号（方便手写表）；整行只有序号没内容则跳过。
    """
    parsed: List[Tuple[int, str]] = []
    auto_no = 0
    for raw in str(text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        match = _STEP_NO_RE.match(line)
        if match:
            auto_no = int(match.group(1))
            content = match.group(2).strip()
        else:
            auto_no += 1
            content = line
        if content:
            parsed.append((auto_no, content))
    return parsed


def _parse_case(case_name: str, steps_text: str, expects_text: str):
    """一行三列 -> (步骤序列, 警告列表)。"""
    warnings: List[str] = []
    parsed_steps = _parse_numbered(steps_text)
    if not parsed_steps:
        return [], warnings

    # 预期按序号归组；同一序号可以有多条（一句播报配多条预期），保持各自独立
    expects_by_no = {}
    for no, text in _parse_numbered(expects_text):
        keywords = _clean_keywords(k.strip() for k in _KEYWORD_SPLIT.split(text))
        if keywords:
            expects_by_no.setdefault(no, []).append(keywords)

    steps: List[VoiceStep] = []
    used_numbers = set()
    for no, text in parsed_steps:
        steps.append(VoiceStep(kind=KIND_PHRASE,
                               text=text[:MAX_CELL_CHARS],
                               delay=DEFAULT_DELAY))
        for keywords in expects_by_no.get(no, []):
            steps.append(VoiceStep(kind=KIND_VERIFY, keywords=list(keywords)))
        if no in expects_by_no:
            used_numbers.add(no)

    for no in expects_by_no:
        if no not in used_numbers:
            warnings.append(f"「{case_name}」序号 {no} 的预期结果找不到对应的操作步骤，已忽略")
    return steps, warnings


def parse_sheet(rows) -> Tuple[List[Tuple[str, List[VoiceStep]]], List[str]]:
    """一个工作表的数据行 -> ([(用例名, 步骤序列)], 警告列表)。

    跳过表头行与空行；没有用例名称、或没有操作步骤的行会被跳过并记一条警告。
    """
    cases: List[Tuple[str, List[VoiceStep]]] = []
    warnings: List[str] = []
    header_pending = True
    for line_no, row in enumerate(rows, start=1):
        cells = [(row[i] if len(row) > i else "") or "" for i in range(3)]
        name, steps_text, expects_text = (str(c).strip() for c in cells)
        if header_pending and name == HEADERS[0]:
            header_pending = False
            continue
        header_pending = False
        if not name and not steps_text and not expects_text:
            continue
        if not name:
            warnings.append(f"第 {line_no} 行没有用例名称，已跳过")
            continue
        steps, case_warnings = _parse_case(name, steps_text, expects_text)
        warnings += case_warnings
        if steps:
            cases.append((name, steps))
        else:
            warnings.append(f"「{name}」没有操作步骤，已跳过")
    return cases, warnings


# ============================================================
# Excel 读写
# ============================================================
def _safe_sheet_name(name: str, used: set) -> str:
    """工作表名合法化：去非法字符、截到 31 字、重名补序号。"""
    clean = _INVALID_SHEET_CHARS.sub("_", (name or "").strip()) or "分组"
    clean = clean[:31]
    candidate, index = clean, 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = clean[:31 - len(suffix)] + suffix
        index += 1
    used.add(candidate)
    return candidate


def write_xlsx(path: str, sheets) -> int:
    """写出 xlsx。sheets = [(分组名, 该组用例的数据行), ...]，一个分组一张表。

    返回写出的工作表数。表名会按 Excel 规则合法化并去重；换行单元格开 wrap_text
    并按行数给个够用的行高，否则打开时只能看到第一行。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    workbook.remove(workbook.active)   # 去掉默认的空表
    used = set()
    wrap = Alignment(wrap_text=True, vertical="top")
    for sheet_name, rows in sheets:
        sheet = workbook.create_sheet(_safe_sheet_name(sheet_name, used))
        sheet.append(list(HEADERS))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            sheet.append([str(c)[:MAX_CELL_CHARS] for c in row])
            for cell in sheet[sheet.max_row]:
                cell.alignment = wrap
            line_count = max((len(str(c).splitlines()) for c in row), default=1)
            sheet.row_dimensions[sheet.max_row].height = \
                max(1, line_count) * ROW_HEIGHT_PER_LINE
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 46
        sheet.column_dimensions["C"].width = 46
    workbook.save(path)
    return len(workbook.sheetnames)


def read_xlsx(path: str) -> List[Tuple[str, List[list]]]:
    """读出 xlsx：返回 [(工作表名, 数据行), ...]，每行已裁成三列。

    表头不在这里剔 —— 统一由 parse_sheet 负责（见那里的注释）。
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    result = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for raw in sheet.iter_rows(values_only=True):
                cells = list(raw[:3]) if raw else []
                while len(cells) < 3:
                    cells.append(None)
                rows.append(["" if c is None else str(c) for c in cells])
            result.append((sheet.title, rows))
    finally:
        workbook.close()
    return result
